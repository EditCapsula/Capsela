import csv, json, re, sys
from collections import Counter
CSV, REVUE, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
VET = {'hauts','pulls_gilets','vestes_blazers','robes','jupes','shorts','manteaux_exterieurs','jeans','pantalons','combinaisons'}
DEJA = re.compile(r'oversize|large|ample|évasé|ajusté|moulant|slim|cintré|côtelé', re.I)  # déjà lu par coupeOf
AJUSTE = [(r'cigarette','pantalon cigarette'),(r'crayon','jupe crayon'),(r'fourreau','robe fourreau'),(r'\bbody\b','body'),(r'legging','legging'),
          (r'brassière','brassière'),(r'bustier','bustier'),(r'bandeau','bandeau'),(r'cycliste|biker','cycliste'),(r'skinny','skinny'),
          (r'col roulé fin|sous-pull','col roulé fin / sous-pull'),(r'maille fine','maille fine')]
AMPLE = [(r'wide','wide leg'),(r'palazzo','palazzo'),(r'relaxed','relaxed'),(r'baggy|boyfriend','baggy / boyfriend'),(r'cargo','cargo'),
         (r'jogger|jogging','jogging'),(r'capuche|hoodie','sweat à capuche'),(r'surchemise','surchemise'),(r'ballon|bouffant','manches ballon / bouffantes'),
         (r'\bcape\b','cape'),(r'kimono|poncho','kimono / poncho'),(r'trapèze|ligne a\b','trapèze / ligne A'),(r'déstructuré','déstructuré'),
         (r'shearling|doudoune','doudoune / shearling')]
def proposer(nom):
    n = nom.lower()
    if re.search(r'flare', n): return '', 'flare : ajusté puis évasé, non tranché'
    for rx, m in AJUSTE:
        if re.search(rx, n): return 'Ajusté', m
    for rx, m in AMPLE:
        if re.search(rx, n): return 'Ample', m
    return '', ''
rows = [x for x in csv.DictReader(open(CSV)) if x['category'] in VET]
crees = [dict(x, id='(créée)') for x in json.load(open(REVUE))['revue'] if x.get('category') in VET]
lignes = []
for x in rows + crees:
    if (x.get('coupe') or '').strip(): continue
    if DEJA.search(x['name']): continue
    c, m = proposer(x['name'])
    lignes.append((x['id'], x['category'], x['name'], x.get('genre',''), x.get('styles',''), x.get('matiere') or '', c, m))
cnt = Counter(l[6] or '(aucune)' for l in lignes)
print('pièces sans coupe à examiner', len(lignes), dict(cnt))
print('dont créées Créatif', sum(1 for l in lignes if l[0]=='(créée)'), Counter(l[6] or '(aucune)' for l in lignes if l[0]=='(créée)'))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); ws = wb.active; ws.title = 'Coupes à valider'
ws.append(['id','catégorie','nom','genre','styles','matière','coupe proposée','motif','validation (Oui / Non / autre valeur)'])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='9C5A44')
for l in sorted(lignes, key=lambda l: (l[6]=='', l[1], l[2])): ws.append(list(l) + [''])
for col, w in zip('ABCDEFGHI', [9,18,46,9,30,22,15,30,30]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
wb.save(OUT)
